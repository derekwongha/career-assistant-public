"""
Unit tests for excel_tracker_writer.py
"""

import os
import shutil
import tempfile
import unittest
from pathlib import Path
import openpyxl

from tracker_template import build_blank_tracker

from excel_tracker_writer import (
    DEFAULT_EXCEL_PATH,
    check_excel_duplicate,
    find_first_blank_placeholder,
    get_sgt_date_str,
    map_work_type,
    resolve_work_type,
    restore_excel_backup,
    write_application_to_excel,
)


class TestExcelTrackerWriter(unittest.TestCase):
    def setUp(self):
        # Build a temporary fictional tracker workbook for unit testing
        self.temp_dir = tempfile.mkdtemp()
        self.temp_excel = Path(self.temp_dir) / "test_tracker.xlsx"
        # Fresh fictional tracker: APP-001..APP-028 used, APP-029 is the first blank placeholder.
        build_blank_tracker(self.temp_excel, prefilled_fictional_rows=28)

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_find_first_blank_placeholder(self):
        wb = openpyxl.load_workbook(self.temp_excel)
        sheet = wb["Applications"]
        row_idx, app_id = find_first_blank_placeholder(sheet)
        self.assertEqual(row_idx, 30)
        self.assertEqual(app_id, "APP-029")

    def test_map_work_type(self):
        self.assertEqual(map_work_type("Hybrid"), "Hybrid")
        self.assertEqual(map_work_type("Full time"), "Full-time")
        self.assertEqual(map_work_type("Contractor"), "Contract")
        self.assertEqual(map_work_type("Not stated"), "Unknown")
        self.assertEqual(map_work_type(None), "Unknown")

    def test_resolve_work_type(self):
        # Test fallback from employment_type ('Full time') over work_mode ('Not stated')
        facts1 = {"employment_type": "Full time", "work_mode": "Not stated"}
        self.assertEqual(resolve_work_type(facts1), "Full-time")

        # Test fallback to work_mode when employment_type is missing/Not stated
        facts2 = {"employment_type": "Not stated", "work_mode": "Hybrid"}
        self.assertEqual(resolve_work_type(facts2), "Hybrid")


    def test_write_application_to_excel_success(self):
        job_data = {
            "job_id": "99999999",
            "confirmed_url": "https://www.jobstreet.com.sg/job/99999999",
            "recommended_priority": "High",
            "analysis": {
                "job": {
                    "company": "Test Company SG",
                    "role": "Full Stack Dev",
                    "location": "Singapore",
                    "work_mode": "Hybrid",
                    "salary": "$4,000",
                },
                "tracker_note": "Great fit for React/Python stack.",
            },
        }

        res = write_application_to_excel(
            excel_path=self.temp_excel,
            job_data=job_data,
            application_url="https://careers.example.com/jobs/123",
            custom_notes="Custom note override",
        )

        self.assertTrue(res["success"])
        self.assertEqual(res["excel_app_id"], "APP-029")
        self.assertEqual(res["row_index"], 30)

        # Inspect written file
        wb = openpyxl.load_workbook(self.temp_excel, data_only=True)
        sheet = wb["Applications"]
        self.assertEqual(sheet.cell(row=30, column=1).value, "APP-029")
        self.assertEqual(sheet.cell(row=30, column=2).value, get_sgt_date_str())
        self.assertEqual(sheet.cell(row=30, column=3).value, "Test Company SG")
        self.assertEqual(sheet.cell(row=30, column=4).value, "Full Stack Dev")
        self.assertEqual(sheet.cell(row=30, column=5).value, "JobStreet")
        self.assertEqual(sheet.cell(row=30, column=6).value, "https://careers.example.com/jobs/123")
        self.assertEqual(sheet.cell(row=30, column=8).value, "Hybrid")
        self.assertEqual(sheet.cell(row=30, column=9).value, "$4,000")

        self.assertEqual(sheet.cell(row=30, column=10).value, "Applied")
        self.assertEqual(sheet.cell(row=30, column=11).value, "High")
        self.assertEqual(sheet.cell(row=30, column=15).value, "Yes")
        self.assertEqual(sheet.cell(row=30, column=17).value, "Custom note override")


    def test_duplicate_prevention(self):
        # Fictional application written into the temporary from-scratch tracker,
        # then written again: the second write must be rejected as a duplicate.
        job_data = {
            "job_id": "FICTIONAL-0001",
            "confirmed_url": "https://example.com/jobs/fictional-0001",
            "recommended_priority": "Medium",
            "analysis": {"job": {"company": "Fictional Company Pte. Ltd.", "role": "Example Developer"}},
        }

        first = write_application_to_excel(self.temp_excel, job_data)
        self.assertEqual(first["excel_app_id"], "APP-029")

        with self.assertRaises(ValueError):
            write_application_to_excel(self.temp_excel, job_data)

        # The rejected second write must not consume another placeholder.
        wb = openpyxl.load_workbook(self.temp_excel, data_only=True)
        sheet = wb["Applications"]
        self.assertIsNone(sheet.cell(row=31, column=3).value)


if __name__ == "__main__":
    unittest.main()
