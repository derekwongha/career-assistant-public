"""
Blank job-application tracker workbook, built from scratch.

Creates the three-sheet layout the Excel writer expects:
- Applications: 17 columns, pre-created placeholder rows APP-001..APP-100
- Dashboard: summary formulas over rows 2:101
- Lists: values used by the dropdown validations

No personal workbook is copied. Document author metadata is cleared.

Usage:
    python 04_Application/tracker_template.py data/job_application_tracker.xlsx
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = [
    "Application ID", "Date Applied", "Company", "Role", "Source", "Job URL",
    "Location", "Work Type", "Salary Range", "Status", "Priority",
    "Contact / Recruiter", "Follow-up Date", "Next Action", "Portfolio Sent",
    "Interview Date", "Outcome / Notes",
]
PLACEHOLDER_ROWS = 100

LISTS = {
    "Source": ["JobStreet", "LinkedIn", "MyCareersFuture", "Company Website", "Recruiter", "Referral", "Other"],
    "Work Type": ["Full-time", "Contract", "Temporary", "Part-time", "Hybrid", "Remote", "On-site", "Unknown"],
    "Status": ["Not Applied", "Applied", "Interview", "Offer", "Rejected", "Withdrawn"],
    "Priority": ["High", "Medium", "Low"],
    "Portfolio Sent": ["Yes", "No"],
}
# Applications column letter for each list-backed field
LIST_COLUMNS = {"Source": "E", "Work Type": "H", "Status": "J", "Priority": "K", "Portfolio Sent": "O"}


def build_blank_tracker(path: Path, prefilled_fictional_rows: int = 0) -> Path:
    """Write a blank tracker to `path`.

    prefilled_fictional_rows fills the first N placeholders with obviously
    fictional applications (used by tests that expect earlier rows in use).
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    apps = wb.active
    apps.title = "Applications"
    dash = wb.create_sheet("Dashboard")
    lists = wb.create_sheet("Lists")

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col, name in enumerate(HEADERS, start=1):
        cell = apps.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
    apps.freeze_panes = "A2"

    last_row = PLACEHOLDER_ROWS + 1
    for i in range(1, PLACEHOLDER_ROWS + 1):
        apps.cell(row=i + 1, column=1, value=f"APP-{i:03d}")

    for i in range(1, prefilled_fictional_rows + 1):
        r = i + 1
        apps.cell(row=r, column=2, value="2026-01-01")
        apps.cell(row=r, column=3, value=f"Fictional Company {i:02d}")
        apps.cell(row=r, column=4, value="Example Role")
        apps.cell(row=r, column=5, value="Other")
        apps.cell(row=r, column=7, value="Singapore")
        apps.cell(row=r, column=8, value="Unknown")
        apps.cell(row=r, column=9, value="Undisclosed")
        apps.cell(row=r, column=10, value="Applied")
        apps.cell(row=r, column=11, value="Medium")
        apps.cell(row=r, column=15, value="Yes")

    # Lists sheet + dropdown validations
    for col, (name, values) in enumerate(LISTS.items(), start=1):
        lists.cell(row=1, column=col, value=name).font = Font(bold=True)
        for r, v in enumerate(values, start=2):
            lists.cell(row=r, column=col, value=v)
        letter = lists.cell(row=1, column=col).column_letter
        dv = DataValidation(
            type="list",
            formula1=f"=Lists!${letter}$2:${letter}${len(values) + 1}",
            allow_blank=True,
        )
        apps.add_data_validation(dv)
        target = LIST_COLUMNS[name]
        dv.add(f"{target}2:{target}{last_row}")

    # Dashboard formulas (rows 2:101 of Applications)
    rng = lambda c: f"Applications!${c}$2:${c}${last_row}"
    dash["A1"] = "Application Tracker Summary"
    dash["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Total applications", f'=COUNTA({rng("C")})'),
        ("Applied", f'=COUNTIF({rng("J")},"Applied")'),
        ("Interview", f'=COUNTIF({rng("J")},"Interview")'),
        ("Offer", f'=COUNTIF({rng("J")},"Offer")'),
        ("Rejected", f'=COUNTIF({rng("J")},"Rejected")'),
        ("High priority", f'=COUNTIF({rng("K")},"High")'),
        ("Medium priority", f'=COUNTIF({rng("K")},"Medium")'),
        ("Low priority", f'=COUNTIF({rng("K")},"Low")'),
    ]
    for r, (label, formula) in enumerate(rows, start=3):
        dash.cell(row=r, column=1, value=label)
        dash.cell(row=r, column=2, value=formula)

    # Clear document author metadata
    props = wb.properties
    props.creator = ""
    props.lastModifiedBy = ""
    props.title = "Job Application Tracker (blank template)"
    props.company = None

    wb.save(path)
    return path


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Usage: python tracker_template.py <output.xlsx>")
    out = build_blank_tracker(Path(sys.argv[1]))
    print(f"Blank tracker written: {out}")
